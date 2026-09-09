import io
import json
import os
import re

import streamlit as st
from docx import Document
from groq import Groq
from pypdf import PdfReader


# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="Resume ATS Analyzer",
    page_icon="📄",
    layout="wide",
)


# ============================================================
# CONFIGURATION
# ============================================================

MODEL = "openai/gpt-oss-20b"


# ============================================================
# ATS JSON SCHEMA
# ============================================================

ATS_SCHEMA = {
    "type": "object",
    "properties": {
        "ats_score": {
            "type": "integer",
            "minimum": 0,
            "maximum": 100
        },
        "score_label": {
            "type": "string"
        },
        "summary": {
            "type": "string"
        },
        "category_scores": {
            "type": "object",
            "properties": {
                "structure_formatting": {
                    "type": "integer",
                    "minimum": 0,
                    "maximum": 20
                },
                "contact_summary": {
                    "type": "integer",
                    "minimum": 0,
                    "maximum": 10
                },
                "skills_keywords": {
                    "type": "integer",
                    "minimum": 0,
                    "maximum": 20
                },
                "experience": {
                    "type": "integer",
                    "minimum": 0,
                    "maximum": 25
                },
                "education_certifications": {
                    "type": "integer",
                    "minimum": 0,
                    "maximum": 10
                },
                "clarity_consistency": {
                    "type": "integer",
                    "minimum": 0,
                    "maximum": 10
                },
                "ats_risk_factors": {
                    "type": "integer",
                    "minimum": 0,
                    "maximum": 5
                }
            },
            "required": [
                "structure_formatting",
                "contact_summary",
                "skills_keywords",
                "experience",
                "education_certifications",
                "clarity_consistency",
                "ats_risk_factors"
            ],
            "additionalProperties": False
        },
        "strengths": {
            "type": "array",
            "items": {
                "type": "string"
            }
        },
        "improvements": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "priority": {
                        "type": "string"
                    },
                    "issue": {
                        "type": "string"
                    },
                    "recommendation": {
                        "type": "string"
                    }
                },
                "required": [
                    "priority",
                    "issue",
                    "recommendation"
                ],
                "additionalProperties": False
            }
        },
        "missing_keywords": {
            "type": "array",
            "items": {
                "type": "string"
            }
        },
        "formatting_risks": {
            "type": "array",
            "items": {
                "type": "string"
            }
        },
        "keyword_alignment": {
            "type": "string"
        },
        "section_feedback": {
            "type": "object",
            "properties": {
                "summary": {
                    "type": "string"
                },
                "skills": {
                    "type": "string"
                },
                "experience": {
                    "type": "string"
                },
                "education": {
                    "type": "string"
                }
            },
            "required": [
                "summary",
                "skills",
                "experience",
                "education"
            ],
            "additionalProperties": False
        }
    },
    "required": [
        "ats_score",
        "score_label",
        "summary",
        "category_scores",
        "strengths",
        "improvements",
        "missing_keywords",
        "formatting_risks",
        "keyword_alignment",
        "section_feedback"
    ],
    "additionalProperties": False
}


# ============================================================
# AI SYSTEM PROMPT
# ============================================================

SYSTEM_PROMPT = """
You are an expert ATS resume analyzer, recruiter, and career consultant.

Your task is to analyze a resume conservatively and realistically.

IMPORTANT RULES:

1. Never invent information.
2. Never assume a skill, qualification, certification, degree,
   technology, job title, responsibility, or achievement unless
   supported by the resume.
3. Do not create fake experience.
4. Give an ATS compatibility estimate, not a guarantee of getting hired.
5. Be objective and constructive.
6. Separate ATS/formatting issues from candidate quality.
7. If a job description is provided, compare the resume against it.
8. Identify important keywords from the job description.
9. Identify keywords that appear to be missing from the resume.
10. Do not recommend adding a keyword if there is no evidence that
    the candidate has relevant experience with it.
11. Do not encourage keyword stuffing.
12. Focus on practical improvements.

ATS SCORE:

The total score must be exactly 100 points.

Structure and formatting = 20 points
Contact information and summary = 10 points
Skills and keywords = 20 points
Professional experience = 25 points
Education and certifications = 10 points
Clarity and consistency = 10 points
ATS risk factors = 5 points

The seven category scores must add up to the final ats_score.

SCORING:

90-100 = Excellent ATS readiness
80-89 = Very Good
70-79 = Good
60-69 = Needs Improvement
Below 60 = Significant Improvement Needed

Consider ATS risks such as:

- Tables
- Multiple columns
- Text boxes
- Graphics
- Images
- Icons replacing text
- Headers and footers containing important information
- Unusual symbols
- Excessive formatting
- Inconsistent dates
- Inconsistent job titles
- Missing contact information
- Poor section headings
- Keyword stuffing
- Long paragraphs
- Unclear employment history
- Missing measurable achievements

When analyzing a job description:

- Identify important keywords.
- Identify relevant keywords already present.
- Identify potentially missing keywords.
- Do not claim that a candidate has a missing skill.
- Recommend adding a keyword only when supported by actual experience.

Return ONLY valid JSON matching the provided schema.
"""


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_pdf_text(file_bytes):
    try:
        reader = PdfReader(io.BytesIO(file_bytes))

        pages = []

        for page in reader.pages:
            text = page.extract_text() or ""
            pages.append(text)

        return "\n".join(pages).strip()

    except Exception as exc:
        raise RuntimeError(
            f"Could not read PDF: {exc}"
        ) from exc


# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

def extract_docx_text(file_bytes):
    try:
        document = Document(io.BytesIO(file_bytes))

        parts = []

        # Paragraphs
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                parts.append(text)

        # Tables
        for table in document.tables:

            for row in table.rows:

                cells = []

                for cell in row.cells:
                    text = cell.text.strip()

                    if text:
                        cells.append(text)

                if cells:
                    parts.append(" | ".join(cells))

        return "\n".join(parts).strip()

    except Exception as exc:
        raise RuntimeError(
            f"Could not read DOCX: {exc}"
        ) from exc


# ============================================================
# TXT TEXT EXTRACTION
# ============================================================

def extract_txt_text(file_bytes):
    try:
        return file_bytes.decode(
            "utf-8",
            errors="ignore"
        ).strip()

    except Exception as exc:
        raise RuntimeError(
            f"Could not read TXT file: {exc}"
        ) from exc


# ============================================================
# GENERAL FILE EXTRACTION
# ============================================================

def extract_resume_text(uploaded_file):

    file_bytes = uploaded_file.getvalue()

    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        return extract_pdf_text(file_bytes)

    elif filename.endswith(".docx"):
        return extract_docx_text(file_bytes)

    elif filename.endswith(".txt"):
        return extract_txt_text(file_bytes)

    else:
        raise RuntimeError(
            "Unsupported file type. Please upload PDF, DOCX, or TXT."
        )


# ============================================================
# CLEAN TEXT
# ============================================================

def clean_text(text):

    if not text:
        return ""

    text = text.replace("\x00", " ")

    text = text.replace(
        "\r\n",
        "\n"
    )

    text = text.replace(
        "\r",
        "\n"
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text
    )

    return text.strip()


# ============================================================
# GET GROQ API KEY
# ============================================================

def get_groq_api_key():

    api_key = None

    try:
        api_key = st.secrets.get(
            "GROQ_API_KEY"
        )
    except Exception:
        pass

    if not api_key:
        api_key = os.getenv(
            "GROQ_API_KEY"
        )

    return api_key


# ============================================================
# GROQ RESUME ANALYSIS
# ============================================================

def analyze_resume(
    resume_text,
    job_description
):

    api_key = get_groq_api_key()

    if not api_key:

        raise RuntimeError(
            "GROQ_API_KEY is missing. "
            "Add GROQ_API_KEY to Streamlit Secrets."
        )

    client = Groq(
        api_key=api_key
    )

    if job_description.strip():

        job_text = job_description.strip()

    else:

        job_text = (
            "No job description was provided. "
            "Analyze the resume for general ATS readiness."
        )

    user_prompt = f"""
RESUME:

{resume_text}


JOB DESCRIPTION:

{job_text}


Analyze this resume according to the system instructions.

Make sure the seven category scores add up exactly to the final ATS score.

Return only valid JSON.
"""

    try:

        response = client.chat.completions.create(

            model=MODEL,

            messages=[
                {
                    "role": "system",
                    "content": SYSTEM_PROMPT
                },
                {
                    "role": "user",
                    "content": user_prompt
                }
            ],

            temperature=0,

            response_format={
                "type": "json_schema",
                "json_schema": {
                    "name": "ats_analysis",
                    "strict": True,
                    "schema": ATS_SCHEMA
                }
            }
        )

    except Exception as exc:

        error_text = str(exc).lower()

        if (
            "429" in error_text
            or "rate limit" in error_text
            or "rate_limit" in error_text
        ):

            raise RuntimeError(
                "Groq rate limit reached. "
                "Please wait and try again later."
            ) from exc

        if (
            "401" in error_text
            or "authentication" in error_text
            or "invalid api key" in error_text
        ):

            raise RuntimeError(
                "Groq API key is invalid. "
                "Check GROQ_API_KEY in Streamlit Secrets."
            ) from exc

        if (
            "model" in error_text
            and (
                "not found" in error_text
                or "does not exist" in error_text
            )
        ):

            raise RuntimeError(
                f"The model {MODEL} is unavailable."
            ) from exc

        raise RuntimeError(
            f"Groq API error: {type(exc).__name__}: {exc}"
        ) from exc

    if not response.choices:

        raise RuntimeError(
            "Groq returned no response."
        )

    output_text = (
        response
        .choices[0]
        .message
        .content
    )

    if not output_text:

        raise RuntimeError(
            "Groq returned an empty response."
        )

    try:

        result = json.loads(
            output_text
        )

    except json.JSONDecodeError as exc:

        raise RuntimeError(
            "Groq returned invalid JSON."
        ) from exc

    return result


# ============================================================
# SCORE LABEL
# ============================================================

def get_score_label(score):

    if score >= 90:
        return "Excellent ATS Readiness"

    if score >= 80:
        return "Very Good ATS Readiness"

    if score >= 70:
        return "Good ATS Readiness"

    if score >= 60:
        return "Needs Improvement"

    return "Significant Improvement Needed"


# ============================================================
# DISPLAY CATEGORY SCORES
# ============================================================

def display_category_scores(category_scores):

    st.subheader(
        "📊 Category Scores"
    )

    categories = [

        (
            "Structure & Formatting",
            "structure_formatting",
            20
        ),

        (
            "Contact & Summary",
            "contact_summary",
            10
        ),

        (
            "Skills & Keywords",
            "skills_keywords",
            20
        ),

        (
            "Experience",
            "experience",
            25
        ),

        (
            "Education & Certifications",
            "education_certifications",
            10
        ),

        (
            "Clarity & Consistency",
            "clarity_consistency",
            10
        ),

        (
            "ATS Risk Factors",
            "ats_risk_factors",
            5
        )
    ]

    cols = st.columns(4)

    for index, (
        label,
        key,
        maximum
    ) in enumerate(categories):

        value = int(
            category_scores.get(
                key,
                0
            )
        )

        with cols[index % 4]:

            st.metric(
                label,
                f"{value}/{maximum}"
            )


# ============================================================
# DISPLAY RESULTS
# ============================================================

def display_results(result):

    score = int(
        result.get(
            "ats_score",
            0
        )
    )

    label = result.get(
        "score_label",
        get_score_label(score)
    )

    # --------------------------------------------------------
    # Overall score
    # --------------------------------------------------------

    st.divider()

    st.subheader(
        "🎯 Overall ATS Score"
    )

    score_col, summary_col = st.columns(
        [1, 2]
    )

    with score_col:

        st.metric(
            "ATS Score",
            f"{score}/100"
        )

        st.progress(
            max(
                0,
                min(
                    score,
                    100
                )
            ) / 100
        )

        st.caption(label)

    with summary_col:

        st.markdown(
            "### Resume Summary"
        )

        st.write(
            result.get(
                "summary",
                "No summary available."
            )
        )

    # --------------------------------------------------------
    # Category scores
    # --------------------------------------------------------

    display_category_scores(
        result.get(
            "category_scores",
            {}
        )
    )

    # --------------------------------------------------------
    # Strengths
    # --------------------------------------------------------

    st.divider()

    st.subheader(
        "✅ Strengths"
    )

    strengths = result.get(
        "strengths",
        []
    )

    if strengths:

        for strength in strengths:

            st.success(
                strength
            )

    else:

        st.info(
            "No specific strengths were identified."
        )

    # --------------------------------------------------------
    # Missing keywords
    # --------------------------------------------------------

    st.subheader(
        "🔑 Missing Keywords"
    )

    missing_keywords = result.get(
        "missing_keywords",
        []
    )

    if missing_keywords:

        st.warning(
            ", ".join(
                missing_keywords
            )
        )

    else:

        st.success(
            "No major missing keywords were identified."
        )

    # --------------------------------------------------------
    # Keyword alignment
    # --------------------------------------------------------

    st.subheader(
        "🔎 Keyword Alignment"
    )

    st.write(
        result.get(
            "keyword_alignment",
            "No keyword alignment analysis available."
        )
    )

    # --------------------------------------------------------
    # Formatting risks
    # --------------------------------------------------------

    st.subheader(
        "⚠️ Formatting & ATS Risks"
    )

    formatting_risks = result.get(
        "formatting_risks",
        []
    )

    if formatting_risks:

        for risk in formatting_risks:

            st.warning(
                risk
            )

    else:

        st.success(
            "No significant ATS formatting risks were identified."
        )

    # --------------------------------------------------------
    # Improvements
    # --------------------------------------------------------

    st.divider()

    st.subheader(
        "🛠️ Recommended Improvements"
    )

    improvements = result.get(
        "improvements",
        []
    )

    if improvements:

        for improvement in improvements:

            priority = improvement.get(
                "priority",
                "Medium"
            )

            issue = improvement.get(
                "issue",
                ""
            )

            recommendation = improvement.get(
                "recommendation",
                ""
            )

            with st.expander(
                f"{priority}: {issue}"
            ):

                st.write(
                    recommendation
                )

    else:

        st.info(
            "No specific improvements were identified."
        )

    # --------------------------------------------------------
    # Section feedback
    # --------------------------------------------------------

    st.divider()

    st.subheader(
        "📝 Section-by-Section Feedback"
    )

    feedback = result.get(
        "section_feedback",
        {}
    )

    tab1, tab2, tab3, tab4 = st.tabs(
        [
            "Summary",
            "Skills",
            "Experience",
            "Education"
        ]
    )

    with tab1:

        st.write(
            feedback.get(
                "summary",
                "No feedback available."
            )
        )

    with tab2:

        st.write(
            feedback.get(
                "skills",
                "No feedback available."
            )
        )

    with tab3:

        st.write(
            feedback.get(
                "experience",
                "No feedback available."
            )
        )

    with tab4:

        st.write(
            feedback.get(
                "education",
                "No feedback available."
            )
        )


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:

    st.title(
        "📄 Resume ATS Analyzer"
    )

    st.markdown(
        """
### Analyze your resume

This application checks:

- ATS compatibility
- Resume structure
- Skills
- Keywords
- Professional experience
- Education
- Certifications
- Formatting risks
- Job description alignment
"""
    )

    st.divider()

    st.markdown(
        "### 🤖 AI Model"
    )

    st.code(
        MODEL,
        language="text"
    )

    st.caption(
        "Powered by Groq"
    )

    st.divider()

    st.markdown(
        "### 🔒 Privacy"
    )

    st.caption(
        "The application does not save uploaded resumes "
        "to its own storage."
    )


# ============================================================
# MAIN PAGE
# ============================================================

st.title(
    "📄 Resume ATS Analyzer"
)

st.markdown(
    """
Upload your resume and receive an ATS-focused analysis.

For better results, optionally paste the job description
you are applying for.
"""
)


# ============================================================
# RESUME UPLOAD
# ============================================================

uploaded_file = st.file_uploader(
    "Upload your resume",
    type=[
        "pdf",
        "docx",
        "txt"
    ],
    help="Supported formats: PDF, DOCX and TXT"
)


# ============================================================
# JOB DESCRIPTION
# ============================================================

job_description = st.text_area(
    "Job Description (Optional)",
    height=220,
    placeholder=(
        "Paste the complete job description here "
        "to analyze keyword alignment..."
    )
)


# ============================================================
# ANALYZE BUTTON
# ============================================================

analyze_button = st.button(
    "🔍 Analyze Resume",
    type="primary",
    use_container_width=True
)


# ============================================================
# RUN ANALYSIS
# ============================================================

if analyze_button:

    if not uploaded_file:

        st.error(
            "Please upload a resume first."
        )

        st.stop()

    try:

        # ----------------------------------------------------
        # Extract resume
        # ----------------------------------------------------

        with st.spinner(
            "Reading your resume..."
        ):

            resume_text = extract_resume_text(
                uploaded_file
            )

            resume_text = clean_text(
                resume_text
            )

        # ----------------------------------------------------
        # Check extracted text
        # ----------------------------------------------------

        if not resume_text:

            st.error(
                "No readable text was found in the uploaded file."
            )

            st.info(
                "If your PDF is a scanned/image-only document, "
                "OCR will be required."
            )

            st.stop()

        # ----------------------------------------------------
        # Limit very large documents
        # ----------------------------------------------------

        MAX_CHARS = 60000

        if len(resume_text) > MAX_CHARS:

            resume_text = resume_text[
                :MAX_CHARS
            ]

            st.warning(
                "The resume was very large. "
                "Only the first 60,000 characters were analyzed."
            )

        # ----------------------------------------------------
        # AI analysis
        # ----------------------------------------------------

        with st.spinner(
            "Analyzing your resume with Groq AI..."
        ):

            result = analyze_resume(
                resume_text,
                job_description
            )

        # ----------------------------------------------------
        # Save result in session
        # ----------------------------------------------------

        st.session_state[
            "ats_result"
        ] = result

        st.success(
            "Resume analysis completed successfully!"
        )

    except Exception as exc:

        st.error(
            f"Analysis failed: {exc}"
        )

        error_text = str(
            exc
        ).lower()

        if "rate limit" in error_text:

            st.warning(
                "The Groq rate limit has been reached. "
                "Please wait before trying again."
            )

        elif "api key" in error_text:

            st.info(
                "Please check GROQ_API_KEY in "
                "Streamlit Cloud → Settings → Secrets."
            )


# ============================================================
# DISPLAY SAVED RESULT
# ============================================================

if (
    "ats_result"
    in st.session_state
):

    display_results(
        st.session_state[
            "ats_result"
        ]
    )


# ============================================================
# FOOTER
# ============================================================

st.divider()

st.caption(
    "Resume ATS Analyzer • AI-assisted resume evaluation • "
    "Always verify recommendations against the actual job requirements."
)
